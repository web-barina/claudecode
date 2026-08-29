# -*- coding: utf-8 -*-
"""Build 中国アプリ事前調査.pptx from research findings."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy

# ---------- palette ----------
NAVY   = RGBColor(0x12, 0x24, 0x38)
INK    = RGBColor(0x1c, 0x2b, 0x35)
MUTED  = RGBColor(0x5c, 0x6f, 0x76)
LINE   = RGBColor(0xd9, 0xe2, 0xe2)
PAPER  = RGBColor(0xf4, 0xf7, 0xf6)
WHITE  = RGBColor(0xff, 0xff, 0xff)
BLUE   = RGBColor(0x1f, 0x5f, 0x96)
BLUE_S = RGBColor(0xe8, 0xf1, 0xfa)
TEAL   = RGBColor(0x12, 0x78, 0x6f)
TEAL_S = RGBColor(0xe6, 0xf4, 0xf1)
AMBER  = RGBColor(0xa8, 0x70, 0x1c)
AMBER_S= RGBColor(0xfa, 0xf1, 0xe0)
RED    = RGBColor(0xb2, 0x3a, 0x2e)
RED_S  = RGBColor(0xfb, 0xea, 0xe8)
ORANGE = RGBColor(0xc6, 0x73, 0x22)
ORANGE_S=RGBColor(0xfc, 0xef, 0xe0)
YELLOW = RGBColor(0x9c, 0x7e, 0x14)
YELLOW_S=RGBColor(0xfa, 0xf5, 0xdf)
GREEN  = RGBColor(0x2c, 0x7a, 0x52)
GREEN_S= RGBColor(0xe7, 0xf4, 0xed)

FONT_HEAD = "Yu Gothic"
FONT_BODY = "Yu Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW = prs.slide_width
SH = prs.slide_height
MARGIN = Inches(0.55)

def add_slide():
    return prs.slides.add_slide(BLANK)

def set_bg(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color

def rect(slide, x, y, w, h, color, line_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp

def rounded(slide, x, y, w, h, color, line_color=None, radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp

def textbox(slide, x, y, w, h, valign=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf

def set_run(r, text, size, color, bold=False, font=FONT_BODY, italic=False):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font

def add_para(tf, text, size, color, bold=False, font=FONT_BODY, space_after=6,
             align=PP_ALIGN.LEFT, first=False, bullet=False, level=0, line_spacing=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    p.level = level
    r = p.add_run()
    prefix = ("・ " if bullet else "")
    set_run(r, prefix + text, size, color, bold, font)
    return p

def eyebrow(slide, text, x=None, y=None, color=TEAL):
    x = MARGIN if x is None else x
    y = Inches(0.42) if y is None else y
    tb, tf = textbox(slide, x, y, Inches(11), Inches(0.35))
    add_para(tf, text.upper(), 11.5, color, bold=True, font="Consolas", first=True)
    return tb

def slide_title(slide, title, y=Inches(0.72), size=27, x=None, w=None):
    x = MARGIN if x is None else x
    w = Inches(12.2) if w is None else w
    tb, tf = textbox(slide, x, y, w, Inches(0.9))
    add_para(tf, title, size, NAVY, bold=True, font=FONT_HEAD, first=True, line_spacing=1.05)
    return tb

def subhead(slide, text, y, size=13.5, color=MUTED, x=None, w=None):
    x = MARGIN if x is None else x
    w = Inches(11.6) if w is None else w
    tb, tf = textbox(slide, x, y, w, Inches(0.55))
    add_para(tf, text, size, color, first=True, line_spacing=1.25)
    return tb

def page_footer(slide, n, section=""):
    tb, tf = textbox(slide, MARGIN, SH - Inches(0.42), Inches(6), Inches(0.3))
    add_para(tf, section, 9.5, MUTED, font="Consolas", first=True)
    tb2, tf2 = textbox(slide, SW - Inches(1.2), SH - Inches(0.42), Inches(0.7), Inches(0.3))
    add_para(tf2, str(n), 9.5, MUTED, font="Consolas", first=True, align=PP_ALIGN.RIGHT)

TAG = {
    "FACT": (BLUE, BLUE_S, "FACT"),
    "HYP": (TEAL, TEAL_S, "整理"),
    "CONFIRM": (AMBER, AMBER_S, "TO CONFIRM"),
}

def tag_chip(slide, x, y, kind, w=Inches(1.5), h=Inches(0.32)):
    color, soft, label = TAG[kind]
    chip = rounded(slide, x, y, w, h, soft, radius=0.5)
    tf = chip.text_frame
    tf.margin_left = Pt(2); tf.margin_right = Pt(2); tf.margin_top=Pt(0); tf.margin_bottom=Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, label, 10.5, color, True, "Consolas")
    return chip

def fact_box(slide, x, y, w, h, kind, body_text, head=None):
    color, soft, label = TAG[kind]
    box = rounded(slide, x, y, w, h, soft, radius=0.045)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14); tf.margin_right = Pt(14); tf.margin_top = Pt(10); tf.margin_bottom = Pt(10)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, head or label, 11, color, True, "Consolas")
    p0.space_after = Pt(6)
    p1 = tf.add_paragraph()
    p1.line_spacing = 1.25
    r1 = p1.add_run()
    set_run(r1, body_text, 13, INK, False, FONT_BODY)
    return box

def plain_box(slide, x, y, w, h, title, body_text, title_color=MUTED):
    box = rounded(slide, x, y, w, h, WHITE, line_color=LINE, radius=0.045)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14); tf.margin_right = Pt(14); tf.margin_top = Pt(10); tf.margin_bottom = Pt(10)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, title, 12, title_color, True, FONT_HEAD)
    p0.space_after = Pt(6)
    p1 = tf.add_paragraph()
    p1.line_spacing = 1.25
    r1 = p1.add_run()
    set_run(r1, body_text, 13, INK, False, FONT_BODY)
    return box

def bullet_box(slide, x, y, w, h, title, items, title_color=NAVY, bg=WHITE, line_color=LINE):
    box = rounded(slide, x, y, w, h, bg, line_color=line_color, radius=0.045)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(14); tf.margin_right = Pt(14); tf.margin_top = Pt(10); tf.margin_bottom = Pt(10)
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    set_run(r0, title, 12.5, title_color, True, FONT_HEAD)
    p0.space_after = Pt(5)
    for it in items:
        p = tf.add_paragraph()
        p.line_spacing = 1.2
        p.space_after = Pt(3)
        r = p.add_run()
        set_run(r, "・ " + it, 11.5, INK, False, FONT_BODY)
    return box

def set_cell(cell, text, size=11, color=INK, bold=False, bg=None, align=PP_ALIGN.LEFT, font=FONT_BODY):
    cell.margin_left = Pt(8); cell.margin_right = Pt(8); cell.margin_top = Pt(5); cell.margin_bottom = Pt(5)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if bg is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.1
    r = p.add_run()
    set_run(r, text, size, color, bold, font)

def add_table(slide, x, y, w, h, headers, rows, col_widths=None, font_size=11, header_bg=RGBColor(0xe9,0xee,0xed)):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gtable.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(w * cw / total))
    for j, htext in enumerate(headers):
        set_cell(table.cell(0, j), htext, size=11, color=NAVY, bold=True, bg=header_bg)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            bg = WHITE if i % 2 == 0 else RGBColor(0xf6,0xf9,0xf8)
            set_cell(table.cell(i+1, j), val, size=font_size, color=INK, bg=bg)
    # remove default table style banding look by keeping our own bg; tone down borders via style id not critical
    return table

def flow_row(slide, x, y, w, h, steps, gap=Inches(0.14)):
    n = len(steps)
    step_w = (w - gap * (n-1)) / n
    cx = x
    for i, s in enumerate(steps):
        box = rounded(slide, cx, y, step_w, h, WHITE, line_color=LINE, radius=0.08)
        top = rect(slide, cx, y, step_w, Inches(0.05), TEAL)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left=Pt(6); tf.margin_right=Pt(6); tf.margin_top=Pt(4); tf.margin_bottom=Pt(4)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 1.05
        r = p.add_run()
        set_run(r, s, 11.5, INK, True, FONT_BODY)
        cx += step_w + gap
        if i < n-1:
            atb, atf = textbox(slide, cx - gap, y + h/2 - Inches(0.15), gap, Inches(0.3))
            ap = atf.paragraphs[0]
            ap.alignment = PP_ALIGN.CENTER
            ar = ap.add_run()
            set_run(ar, "→", 14, MUTED, True)

# ============================================================
# 1. TITLE SLIDE
# ============================================================
s = add_slide()
set_bg(s, NAVY)
rect(s, 0, 0, SW, Inches(0.09), TEAL)
tb, tf = textbox(s, MARGIN, Inches(0.9), Inches(10), Inches(0.4))
add_para(tf, "PRE-MEETING RESEARCH / COMPILED 2026-08-26", 13, RGBColor(0x8f,0xd6,0xcb), bold=True, font="Consolas", first=True)
tb, tf = textbox(s, MARGIN, Inches(1.7), Inches(11.6), Inches(2.4))
add_para(tf, "中国向けアプリ開発", 42, WHITE, bold=True, font=FONT_HEAD, first=True, line_spacing=1.15)
add_para(tf, "事前調査ファインディングス", 42, WHITE, bold=True, font=FONT_HEAD, line_spacing=1.15)
tb, tf = textbox(s, MARGIN, Inches(3.85), Inches(10.5), Inches(1.1))
add_para(tf, "「中国向けアプリ開発 社内ブレインストーミング」で整理した調査項目を、実際にWeb調査してFACTとして確認できるところまで確認した結果。", 15.5, RGBColor(0xcd,0xda,0xdb), first=True, line_spacing=1.4)
add_para(tf, "明後日のクライアント初回打ち合わせ用の内部資料。", 15.5, RGBColor(0xcd,0xda,0xdb), line_spacing=1.4)

legend_y = Inches(5.35)
labels = [("FACT — 公式情報で確認", BLUE, BLUE_S), ("整理 — 二次情報＋論理整理", TEAL, TEAL_S), ("TO CONFIRM — 面談/専門家確認", AMBER, AMBER_S)]
cx = MARGIN
for label, c, soft in labels:
    w = Inches(3.55)
    chip = rounded(s, cx, legend_y, w, Inches(0.5), RGBColor(0x1a,0x30,0x46), radius=0.5)
    chip.line.color.rgb = c
    chip.line.width = Pt(1)
    tf = chip.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, label, 11.5, WHITE, True)
    cx += w + Inches(0.2)

tb, tf = textbox(s, MARGIN, SH - Inches(0.7), Inches(8), Inches(0.4))
add_para(tf, "中国コインランドリーアプリ開発案件 ／ 社内検討資料", 11.5, RGBColor(0x7f,0x92,0x98), font="Consolas", first=True)

print("title slide done")
